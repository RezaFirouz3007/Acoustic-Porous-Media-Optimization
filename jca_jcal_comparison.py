######################################################################
############## Reminder change the thickness d #######################
######################################################################
# STEP 0 — What we are building
# Comprehensive JCA vs JCAL comparison with multi-parameter sensitivity and full data tracking
######################################################################
# STEP 1 — Import tools
import numpy as np
import matplotlib.pyplot as plt
from scipy.optimize import minimize
import matplotlib
import os, sys, glob
import pandas as pd
from openpyxl import load_workbook
from openpyxl.drawing.image import Image
from openpyxl.utils.dataframe import dataframe_to_rows

matplotlib.use('Agg')
######################################################################
# STEP 2 — Define frequency axis
f_init = np.linspace(100, 4000, 500)
omega_init = 2 * np.pi * f_init
######################################################################
# STEP 3 — Air constants
rho0 = 1.21; c0 = 343; Z0 = rho0 * c0; eta = 1.84e-5
gamma = 1.4; P0 = 101325; Pr = 0.71
kappa = 0.026; Cp = 1006 
j = 1j
######################################################################
# STEP 4 — Foam parameters (Initial Guess)
phi_init = 0.95; sigma_init = 12000; alpha_inf_init = 1.1
Lambda_init = 100e-6; Lambda_p_init = 200e-6; k0p_init = 4e-9; d = 0.075

x0 = [phi_init, sigma_init, alpha_inf_init, Lambda_init, Lambda_p_init, k0p_init]

# Optimization Bounds
bounds = [
    (0.9, 0.999),     # porosity
    (5000, 40000),    # flow resistivity
    (1.0, 2.0),       # tortuosity
    (10e-6, 400e-6),  # viscous length
    (50e-6, 800e-6),  # thermal length
    (1e-10, 5e-8)     # k0p
]
######################################################################
# STEP 5 & 6 — Effective Properties Functions

def get_rho_eff(w, phi, sigma, alpha_inf, Lambda):
    term = np.sqrt(1 + (4 * j * alpha_inf**2 * eta * rho0 * w) / (sigma**2 * phi**2 * Lambda**2))
    return (rho0 * alpha_inf / phi) * (1 + (sigma * phi) / (j * w * rho0 * alpha_inf) * term)

def get_K_eff_jca(w, phi, Lambda_p):
    thermal_term = np.sqrt(1 + (j * w * rho0 * Pr * Lambda_p**2) / (16 * eta))
    return (gamma * P0 / phi) / (gamma - (gamma - 1) / (1 + (8 * eta) / (j * w * rho0 * Pr * Lambda_p**2) * thermal_term))

def get_K_eff_jcal(w, phi, Lambda_p, k0p):
    thermal_term = np.sqrt(1 + (4j * k0p**2 * Cp * rho0 * w) / (kappa * Lambda_p**2 * phi**2))
    denom_internal = 1 - (j * phi * kappa / (k0p * Cp * rho0 * w)) * thermal_term
    return (gamma * P0 / phi) / (gamma - (gamma - 1) / denom_internal)

def calculate_abs(params, f_vec, model_type='jcal'):
    phi, sigma, alpha_inf, L, Lp, k0p_val = params
    w = 2 * np.pi * f_vec
    rho_e = get_rho_eff(w, phi, sigma, alpha_inf, L)
    K_e = get_K_eff_jcal(w, phi, Lp, k0p_val) if model_type == 'jcal' else get_K_eff_jca(w, phi, Lp)
    Zc = np.sqrt(rho_e * K_e); k_prop = w * np.sqrt(rho_e / K_e)
    Zs = -1j * Zc / np.tan(k_prop * d)
    return 1 - np.abs((Zs - Z0) / (Zs + Z0))**2

######################################################################
# STEP 13 — Load experimental data
xlsx_files = glob.glob("*.xlsx")
if not xlsx_files: sys.exit("No .xlsx files found.")
df = pd.read_excel(xlsx_files[0])
f_meas = df["frequency"].to_numpy()
alpha_meas = df["absorption"].to_numpy()

####################################################################
# STEP 16 — Optimization
def error_func(p):
    return np.mean((alpha_meas - calculate_abs(p, f_meas, 'jcal'))**2)

res = minimize(error_func, x0, bounds=bounds, method='L-BFGS-B')
best_p = res.x

####################################################################
# STEP 17 & 18 — Evaluation & Metrics
alpha_jca = calculate_abs(best_p, f_meas, 'jca')
alpha_jcal = calculate_abs(best_p, f_meas, 'jcal')

def get_metrics(meas, model):
    mse = np.mean((meas - model)**2)
    rmse = np.sqrt(mse)
    r2 = 1 - (np.sum((meas - model)**2) / np.sum((meas - np.mean(meas))**2))
    return mse, rmse, r2

mse_jcal, rmse_jcal, r2_jcal = get_metrics(alpha_meas, alpha_jcal)
mse_jca, rmse_jca, r2_jca = get_metrics(alpha_meas, alpha_jca)

# Plot Comparison
plt.figure(figsize=(10, 6))
plt.semilogx(f_meas, alpha_meas, 'r', label="Measured")
plt.semilogx(f_meas, alpha_jca, 'b--', label="JCA Model")
plt.semilogx(f_meas, alpha_jcal, 'k-', label="JCAL Model")
plt.legend(); plt.grid(True, which="both"); plt.title("JCA vs JCAL Fit Comparison")
plt.savefig("fit_comparison.png"); plt.close()

#######################################################################
# STEP 19 — Sensitivity Analysis (All Parameters)
labels = ["Porosity", "Flow Resistivity", "Tortuosity", "Viscous Length", "Thermal Length", "k0p"]
fig, axes = plt.subplots(3, 2, figsize=(15, 12))
axes = axes.flatten()

for i, label in enumerate(labels):
    axes[i].semilogx(f_meas, alpha_meas, 'r', label="Measured", linewidth=2)
    for factor in [0.9, 1.1]:
        p_sens = list(best_p)
        p_sens[i] *= factor
        # Note: k0p variations only affect JCAL
        alpha_sens = calculate_abs(p_sens, f_meas, 'jcal')
        axes[i].semilogx(f_meas, alpha_sens, '--', label=f"{label} {factor}x")
    axes[i].set_title(f"Sensitivity: {label}")
    axes[i].legend(fontsize='small'); axes[i].grid(True)

plt.tight_layout()
plt.savefig("full_sensitivity_analysis.png")
plt.close()

#######################################################################
# STEP 20 — Expanded Excel Export
filename = xlsx_files[0]
book = load_workbook(filename)
if "Comparison_Results" in book.sheetnames: del book["Comparison_Results"]
sheet = book.create_sheet("Comparison_Results")

# 1. Main Data Table
df_export = pd.DataFrame({
    "Frequency (Hz)": f_meas, 
    "Measured": alpha_meas, 
    "JCA_Alpha": alpha_jca, 
    "JCAL_Alpha": alpha_jcal
})
for row in dataframe_to_rows(df_export, index=False, header=True): 
    sheet.append(row)

# 2. Parameters, Metrics, Guesses, and Bounds
sheet["F2"] = "Parameter Name"; sheet["G2"] = "Fitted (JCAL)"; sheet["H2"] = "Initial Guess"; sheet["I2"] = "Bounds (Min, Max)"

for i, (l, b_p, i_p, bnd) in enumerate(zip(labels, best_p, x0, bounds), 3):
    sheet[f"F{i}"] = l
    sheet[f"G{i}"] = float(b_p)
    sheet[f"H{i}"] = float(i_p)
    sheet[f"I{i}"] = str(bnd)

# 3. Fit Quality Metrics
sheet["F10"] = "Fit Metric"; sheet["G10"] = "JCAL Value"; sheet["H10"] = "JCA Value"
sheet["F11"] = "MSE"; sheet["G11"] = float(mse_jcal); sheet["H11"] = float(mse_jca)
sheet["F12"] = "RMSE"; sheet["G12"] = float(rmse_jcal); sheet["H12"] = float(rmse_jca)
sheet["F13"] = "R2"; sheet["G13"] = float(r2_jcal); sheet["H13"] = float(r2_jca)

# Insert Comparison Image
img = Image("fit_comparison.png")
sheet.add_image(img, "K2")

# Save workbook
book.save(filename)
print(f"✅ Final analysis complete. Sensitivity saved as 'full_sensitivity_analysis.png'.")
print(f"✅ Comparison data, initial guesses, and bounds saved to {filename}.")